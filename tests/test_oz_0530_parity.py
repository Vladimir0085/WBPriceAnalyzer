"""Interface parity with OZ Price Analyzer 0.5.30 (WB 0.2.15), synthetic data only."""
from __future__ import annotations

import tempfile
import tkinter as tk
import unittest
from datetime import date
from pathlib import Path
from tkinter import ttk
from types import SimpleNamespace
from unittest.mock import patch

from openpyxl import load_workbook

from wb_app.exporter import export_calculation
from wb_app.models import Product, ProductResult, RunCalculation
from wb_app.report_exports import _find_button
from wb_app.service import AppService
from wb_app.ui import OVERVIEW_COLUMN_SPECS, SCENARIO_COLUMN_SPECS
from wb_app.ui_containers import wrapped_positions
from wb_app.ui_layout import DisplayWBPriceAnalyzerApp, density_style_options, ui_scale_status


def _calculation(*products: ProductResult, unallocated: float = 0.0) -> RunCalculation:
    return RunCalculation(
        run_id=None,
        period_start=date(2026, 7, 1),
        period_end=date(2026, 7, 7),
        tax_rate=0.06,
        products=list(products),
        unallocated_total=unallocated,
        unallocated={"Хранение": (1, unallocated)} if unallocated else {},
        accrual_stats={},
    )


def _pots(units: int = 2) -> ProductResult:
    return ProductResult(
        "A-10", "Товар A", 60, 40, category="Горшки",
        units=units, revenue_no_points=500, financial_result=300,
    )


def _bows() -> ProductResult:
    return ProductResult(
        "B-20", "Товар B", 40, 10, category="Банты",
        units=5, revenue_no_points=900, financial_result=600,
    )


class ParityHelperTests(unittest.TestCase):
    def test_table_headings_match_oz(self) -> None:
        overview = {column: heading for column, heading, _width in OVERVIEW_COLUMN_SPECS}
        scenario = {column: heading for column, heading, _width in SCENARIO_COLUMN_SPECS}
        self.assertEqual(overview["profit_unit"], "Финрезультат WB на ед.")
        self.assertEqual(overview["profit_total"], "Финрезультат WB до с/с и налога")
        self.assertEqual(scenario["profit"], "Прибыль до себестоимости")

    def test_density_follows_oz_and_is_clamped(self) -> None:
        self.assertEqual(density_style_options(1.0)["Treeview"]["rowheight"], 30)
        self.assertEqual(density_style_options(0.9)["Treeview"]["rowheight"], 27)
        self.assertEqual(density_style_options(0.8)["Treeview"]["rowheight"], 24)
        self.assertEqual(density_style_options(0.8)["TButton"]["padding"], (11, 6))
        # DPI emulation above 100% must not inflate paddings.
        self.assertEqual(density_style_options(2.0), density_style_options(1.0))
        self.assertEqual(density_style_options(0.5), density_style_options(0.8))

    def test_scale_status_matches_oz(self) -> None:
        self.assertEqual(ui_scale_status("auto", 0.9), "Масштаб интерфейса: 90% (авто)")
        self.assertEqual(ui_scale_status("0.8", 0.8), "Масштаб интерфейса: 80%")

    def test_right_aligned_rows_end_at_the_toolbar_edge(self) -> None:
        sizes = [(300, 30), (200, 40), (150, 30)]
        positions, height = wrapped_positions(sizes, 1000, align_right=True)
        self.assertEqual(positions, [(334, 0), (642, 0), (850, 0)])
        self.assertEqual(height, 40)
        positions, height = wrapped_positions(sizes, 520, align_right=True)
        # First row: 300 + 8 + 200 = 508; second row keeps only the last group.
        self.assertEqual(positions, [(12, 0), (320, 0), (370, 44)])
        self.assertEqual(height, 74)
        self.assertEqual(wrapped_positions(sizes, 520)[0][2], (0, 44))

    def test_xlsx_uses_activity_profit_names(self) -> None:
        calculation = _calculation(_pots(), unallocated=-120.0)
        with tempfile.TemporaryDirectory() as directory:
            path = export_calculation(calculation, Path(directory) / "report.xlsx")
            workbook = load_workbook(path)
            sheet = workbook["Итог"]
            labels = [sheet.cell(row, 1).value for row in range(1, sheet.max_row + 1)]
            j2 = sheet["J2"].value
            workbook.close()
        self.assertEqual(j2, "Чистая прибыль от деятельности")
        self.assertIn("Итого от деятельности (с нераспределенными)", labels)
        self.assertNotIn("Итого с нераспределенными", labels)


class _TkAppTestCase(unittest.TestCase):
    """Run against real widgets on Windows or with a display (e.g. xvfb-run)."""

    @classmethod
    def setUpClass(cls) -> None:
        try:
            probe = tk.Tk()
        except tk.TclError as exc:
            raise unittest.SkipTest(f"Tk display is unavailable: {exc}") from exc
        probe.destroy()

    def setUp(self) -> None:
        from wb_app.report_totals import ReportTotalsWBPriceAnalyzerApp

        self.directory = tempfile.TemporaryDirectory()
        self.addCleanup(self.directory.cleanup)
        self.callback_errors: list[object] = []
        callback_patch = patch(
            "tkinter.Tk.report_callback_exception",
            side_effect=lambda *error: self.callback_errors.append(error),
        )
        callback_patch.start()
        self.addCleanup(callback_patch.stop)
        initialize_tk = tk.Tk.__init__

        def initialize_at_test_dpi(root, *args, **kwargs):
            initialize_tk(root, *args, **kwargs)
            root.tk.call("tk", "scaling", 96 / 72)

        # A small CI desktop must not maximize the window the tests resize.
        with patch("tkinter.Tk.__init__", initialize_at_test_dpi), patch.object(
            DisplayWBPriceAnalyzerApp, "_maximize_on_small_screen", lambda self: None
        ):
            self.app = ReportTotalsWBPriceAnalyzerApp(AppService(Path(self.directory.name)))
        self.addCleanup(self.app.destroy)
        self.app.minsize(1, 1)
        self.app.maxsize(2200, 1400)
        self.app._base_tk_scaling = 96 / 72
        self.app._apply_ui_scale(1.0)

    def settle(self) -> None:
        for _ in range(4):
            self.app.update_idletasks()
            self.app.update()
        self.assertEqual(self.callback_errors, [])

    def show(self, calculation: RunCalculation, run_ids: set[int] | None = None) -> None:
        self.app.current_calculation = calculation
        self.app.overview_calculation = calculation
        self.app.overview_run_ids = set(run_ids or {1})
        self.app.overview_file_count = 1
        self.app._populate_overview()
        self.app._populate_scenario()

    def label_texts(self, root: tk.Misc) -> list[str]:
        texts = []
        for child in root.winfo_children():
            try:
                texts.append(str(child.cget("text")))
            except tk.TclError:
                pass
            texts.extend(self.label_texts(child))
        return texts


class OverviewParityTests(_TkAppTestCase):
    def test_report_cards_use_oz_names(self) -> None:
        texts = self.label_texts(self.app.kpi_frame)
        self.assertIn("Чистая прибыль товаров", texts)
        self.assertIn("Чистая прибыль от деятельности", texts)
        self.assertNotIn("Итог с нераспределёнными", texts)
        self.show(_calculation(_pots(), _bows(), unallocated=-120.0))
        # Products: (300 − 200 − 30) + (600 − 250 − 54) = 366 ₽; activity: 366 − 120.
        self.assertEqual(self.app.kpi_vars["net_profit"].get(), "366.00 ₽")
        self.assertEqual(self.app.kpi_vars["report_total"].get(), "246.00 ₽")

    def test_category_cards_are_not_clipped_on_a_full_hd_window(self) -> None:
        self.show(_calculation(_pots(), _bows()))
        for size in ("1540x920", "1500x900", "1280x720"):
            with self.subTest(size=size):
                self.app.geometry(f"{size}+0+0")
                self.settle()
                summary = self.app._table_layouts[self.app.overview_tab].upper
                reachable = summary.canvas.winfo_height() >= summary.content.winfo_reqheight()
                self.assertTrue(reachable or summary.yscroll.winfo_ismapped(), size)
                if not summary.yscroll.winfo_ismapped():
                    self.assertEqual(summary.canvas.winfo_height(), summary.winfo_height())

    def test_empty_selection_with_exclude_keeps_every_row(self) -> None:
        self.show(_calculation(_pots(), _bows()))
        self.app._category_exclude_vars["overview"].set(True)
        self.app._on_category_mode_changed("overview")
        self.assertEqual(len(self.app.overview_tree.get_children("")), 2)
        self.assertEqual(
            self.app.category_summary_title_var.get(), "Итоги по фильтру: Все товары · 2 позиции"
        )

    def test_selected_and_excluded_categories_update_rows_and_title(self) -> None:
        self.show(_calculation(_pots(), _bows()))
        self.app._category_selected["overview"] = {"Горшки"}
        self.app._on_category_mode_changed("overview")
        self.assertEqual(self.app.overview_tree.get_children(""), ("A-10",))
        self.assertEqual(
            self.app.category_summary_title_var.get(),
            "Итоги по фильтру: категории: Горшки · 1 позиция",
        )
        self.assertEqual(self.app.category_kpi_vars["revenue"].get(), "500.00 ₽")
        self.app._category_exclude_vars["overview"].set(True)
        self.app._on_category_mode_changed("overview")
        self.assertEqual(self.app.overview_tree.get_children(""), ("B-20",))
        self.assertEqual(
            self.app.category_summary_title_var.get(),
            "Итоги по фильтру: кроме категорий: Горшки · 1 позиция",
        )

    def test_category_dialog_matches_oz(self) -> None:
        from wb_app.ui_categories import CategorySelectionDialog

        dialog = CategorySelectionDialog(
            self.app, ["Банты", "Горшки"], set(), title="Категории — Обзор"
        )
        self.assertEqual(dialog.title(), "Категории — Обзор")
        self.assertEqual(dialog.summary_var.get(), "Все категории")
        texts = self.label_texts(dialog)
        self.assertIn("Выберите одну или несколько категорий", texts)
        self.assertIn("Снять все", texts)
        dialog._select_all()
        self.assertEqual(dialog.summary_var.get(), "Выбрано: 2")
        dialog._clear()
        dialog._apply()
        self.assertFalse(dialog.cancelled)
        self.assertEqual(dialog.result, set())

    def test_no_categories_message_matches_oz(self) -> None:
        with patch("wb_app.ui_categories.messagebox.showinfo") as info:
            self.app._choose_categories("overview")
        info.assert_called_once_with(
            "Фильтр категорий", "В текущем отчете нет категорий для выбора.", parent=self.app
        )


class ScenarioParityTests(_TkAppTestCase):
    def test_cards_and_categories_follow_every_selected_report(self) -> None:
        combined = _calculation(_pots(), _bows())
        self.show(combined, run_ids={1, 2})
        # The report open in the top list contains only one of the products.
        self.app.current_calculation = _calculation(_pots())
        self.app._populate_scenario()
        self.assertEqual(self.app.scenario_tree.get_children(""), ("A-10", "B-20"))
        self.assertEqual(self.app.scenario_kpi_vars["current_revenue"].get(), "1 400.00 ₽")
        self.assertEqual(self.app._available_categories("scenario"), ["Банты", "Горшки"])
        self.assertIn("2 отчета · цены только для 1 отчета", self.app.scenario_count_var.get())

    def test_cards_sit_under_the_table_and_hide_in_fullscreen(self) -> None:
        self.show(_calculation(_pots()))
        layout = self.app._table_layouts[self.app.scenario_tab]
        self.assertIs(self.app.scenario_kpi_frame.master, layout.table)
        self.app.geometry("1540x920+0+0")
        self.app.notebook.select(self.app.scenario_tab)
        self.settle()
        self.assertGreater(
            self.app.scenario_kpi_frame.winfo_rooty(), self.app.scenario_tree.winfo_rooty()
        )
        controller = self.app._table_modes["scenario"]
        controller.expand()
        self.settle()
        self.assertFalse(self.app.scenario_kpi_frame.winfo_ismapped())
        self.assertEqual(str(controller.detach_button.cget("text")), "Открыть отдельно")
        # Esc restores the table even after switching to another tab.
        self.app.notebook.select(self.app.overview_tab)
        self.assertEqual(self.app._escape_table_mode(), "break")
        self.settle()
        self.assertFalse(controller.fullscreen)
        self.assertEqual(str(controller.detach_button.cget("text")), "↗ Отдельно")
        self.app.notebook.select(self.app.scenario_tab)
        self.settle()
        self.assertTrue(self.app.scenario_kpi_frame.winfo_ismapped())

    def test_column_settings_status_matches_oz(self) -> None:
        confirmed = SimpleNamespace(
            confirmed=True, preferences=list(self.app.scenario_column_preferences)
        )
        with patch("wb_app.column_settings.ColumnSettingsDialog", return_value=confirmed), patch.object(
            self.app, "wait_window"
        ):
            self.app.open_scenario_column_settings()
            self.assertEqual(
                self.app.status_var.get(), "Настройка столбцов сценария цены сохранена"
            )
            self.app.open_overview_column_settings()
        self.assertEqual(self.app.status_var.get(), "Настройка столбцов обзора сохранена")


class CatalogAndWindowParityTests(_TkAppTestCase):
    def test_catalog_tab_matches_oz(self) -> None:
        delete = _find_button(self.app.catalog_tab, "Удалить")
        self.assertEqual(str(delete.cget("style")), "Danger.TButton")
        self.assertIn(
            "Справочник влияет только на будущие расчеты. Сохраненные отчеты хранят "
            "исторический снимок себестоимости.",
            self.label_texts(self.app.catalog_tab),
        )

    def test_skipped_rows_button_only_with_skipped_rows(self) -> None:
        self.assertEqual(self.app.catalog_warning_frame.winfo_manager(), "")
        self.app.db.set_setting(
            "cost_catalog_warnings", "Строка 5, артикул Z-1: не указана полная себестоимость"
        )
        self.app._refresh_cost_catalog_warning()
        self.assertEqual(self.app.catalog_warning_frame.winfo_manager(), "grid")
        self.app.db.set_setting("cost_catalog_warnings", "")
        self.app._refresh_cost_catalog_warning()
        self.assertEqual(self.app.catalog_warning_frame.winfo_manager(), "")

    def test_delete_texts_match_oz(self) -> None:
        self.app.db.save_product(Product("K-1", "Кашпо", 30, 10))
        self.app.refresh_products()
        self.app.products_tree.selection_set("K-1")
        with patch("wb_app.ui_catalog.messagebox.askyesno", return_value=True) as ask:
            self.app.delete_selected_product()
        title, question = ask.call_args.args
        self.assertEqual(title, "Удалить позицию из справочника?")
        self.assertIn("Удалить «Кашпо» (артикул K-1) из текущего справочника себестоимости?", question)
        self.assertIn("его потребуется добавить заново", question)
        self.assertEqual(self.app.status_var.get(), "Из справочника удалена позиция: K-1")
        self.assertNotIn("K-1", self.app.db.product_map())
        with patch("wb_app.ui_catalog.messagebox.showinfo") as info:
            self.app.delete_selected_product()
        info.assert_called_once_with(
            "Удалить товар", "Выберите позицию в справочнике себестоимости.", parent=self.app
        )

    def test_scale_density_survives_theme_change(self) -> None:
        style = ttk.Style(self.app)
        self.assertEqual(int(style.lookup("Treeview", "rowheight")), 30)
        self.app.ui_scale_var.set("80%")
        self.app._scale_changed()
        self.assertEqual(self.app.status_var.get(), "Масштаб интерфейса: 80%")
        self.assertEqual(int(style.lookup("Treeview", "rowheight")), 24)
        self.assertFalse(self.app.header_subtitle.winfo_ismapped())
        self.app._preview_theme()
        self.assertEqual(int(style.lookup("Treeview", "rowheight")), 24)
        self.app.ui_scale_var.set("100%")
        self.app._scale_changed()
        self.settle()
        self.assertTrue(self.app.header_subtitle.winfo_ismapped())

    def test_header_buttons_are_not_clipped(self) -> None:
        toolbar = self.app.header_actions
        groups = toolbar.groups
        about = groups[-1].winfo_children()[0]
        self.assertEqual(str(about.cget("text")), "О программе")
        # One row or a wrap depends on the font (Segoe UI on Windows, DejaVu on
        # Linux), so both are valid; only clipping and overflow are errors.
        for size in ("1540x920", "1280x720"):
            with self.subTest(size=size):
                self.app.geometry(f"{size}+0+0")
                self.settle()
                left = self.app.winfo_rootx()
                right = left + self.app.winfo_width()
                bottom = self.app.winfo_rooty() + self.app.winfo_height()
                title_right = (
                    self.app.header_subtitle.winfo_rootx() + self.app.header_subtitle.winfo_width()
                )
                for group in groups:
                    self.assertTrue(group.winfo_ismapped())
                    self.assertGreaterEqual(group.winfo_x(), 0)
                    self.assertGreaterEqual(group.winfo_y(), 0)
                    self.assertLessEqual(group.winfo_x() + group.winfo_width(), toolbar.winfo_width())
                    self.assertLessEqual(group.winfo_y() + group.winfo_height(), toolbar.winfo_height())
                    self.assertGreaterEqual(group.winfo_rootx(), max(left, title_right))
                    self.assertLessEqual(group.winfo_rootx() + group.winfo_width(), right)
                    self.assertLessEqual(group.winfo_rooty() + group.winfo_height(), bottom)
                    for control in group.winfo_children():
                        self.assertEqual(control.winfo_width(), control.winfo_reqwidth())
                        self.assertEqual(control.winfo_height(), control.winfo_reqheight())
                for index, first in enumerate(groups):
                    for second in groups[index + 1:]:
                        self.assertTrue(
                            first.winfo_x() + first.winfo_width() <= second.winfo_x()
                            or second.winfo_x() + second.winfo_width() <= first.winfo_x()
                            or first.winfo_y() + first.winfo_height() <= second.winfo_y()
                            or second.winfo_y() + second.winfo_height() <= first.winfo_y(),
                            (size, first.winfo_geometry(), second.winfo_geometry()),
                        )


if __name__ == "__main__":
    unittest.main()
