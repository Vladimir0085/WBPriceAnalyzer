"""Эталон расчётов на синтетических отчётах Wildberries.

Эталон ``fixtures/calculation_etalon.json`` снят на версии 0.2.15 до
интерфейсных изменений 0.2.16. Изменения интерфейса, подписей, подсказок и
справки не должны менять ни одной копейки: выручку MAIN и BUYOUT, продажи,
налог, чистую прибыль товаров, чистую прибыль от деятельности, доходность,
доли от выручки, сценарий цены и ячейки экспорта XLSX.

Налог по методике WB: база = Выручка MAIN + Выручка BUYOUT, налог с
нераспределённых операций не начисляется — это проверяется отдельно.

Перезаписывать эталон можно только при осознанном изменении расчёта:
``python -m tests.test_calculation_etalon --write``.
"""

from __future__ import annotations

import json
import os
import sys
import tempfile
import unittest
from pathlib import Path
from unittest.mock import patch

from openpyxl import Workbook, load_workbook

from wb_app.aggregation import aggregate_calculations
from wb_app.calculator import calculate_scenario
from wb_app.exporter import export_calculation
from wb_app.models import Product, RunCalculation
from wb_app.service import AppService


FIXTURE = Path(__file__).resolve().parent / "fixtures" / "calculation_etalon.json"
MONEY_TOLERANCE = 1e-6
RATIO_TOLERANCE = 1e-9
TAX_RATE = "0.06"

HEADERS = [
    "Предмет", "Код номенклатуры", "Артикул поставщика", "Название", "Тип документа",
    "Обоснование для оплаты", "Дата заказа покупателем", "Дата продажи", "Кол-во",
    "Цена розничная", "Вайлдберриз реализовал Товар (Пр)",
    "Возмещение за выдачу и возврат товаров на ПВЗ",
    "Компенсация платёжных услуг/Комиссия за интеграцию платёжных сервисов",
    "Вознаграждение Вайлдберриз (ВВ), без НДС", "НДС с Вознаграждения Вайлдберриз",
    "К перечислению Продавцу за реализованный Товар", "Услуги по доставке товара покупателю",
    "Общая сумма штрафов", "Корректировка Вознаграждения Вайлдберриз (ВВ)",
    "Виды логистики, штрафов и корректировок ВВ", "Страна", "Srid",
    "Возмещение издержек по перевозке/по складским операциям с товаром", "Хранение",
    "Удержания", "Операции на приемке", "Компенсация скидки по программе лояльности",
    "Стоимость участия в программе лояльности", "Сумма баллов, удержанных по программе лояльности",
    "Разовое изменение срока перечисления денежных средств", "Коэффициент логистики",
]
COLUMN = {name: index for index, name in enumerate(HEADERS)}
FIELD_COLUMNS = {
    "retail": "Цена розничная",
    "realized": "Вайлдберриз реализовал Товар (Пр)",
    "pvz": "Возмещение за выдачу и возврат товаров на ПВЗ",
    "acquiring": "Компенсация платёжных услуг/Комиссия за интеграцию платёжных сервисов",
    "commission": "Вознаграждение Вайлдберриз (ВВ), без НДС",
    "commission_vat": "НДС с Вознаграждения Вайлдберриз",
    "payout": "К перечислению Продавцу за реализованный Товар",
    "logistics": "Услуги по доставке товара покупателю",
    "penalty": "Общая сумма штрафов",
    "storage": "Хранение",
    "deductions": "Удержания",
    "acceptance": "Операции на приемке",
    "loyalty_compensation": "Компенсация скидки по программе лояльности",
    "loyalty_fee": "Стоимость участия в программе лояльности",
    "loyalty_points": "Сумма баллов, удержанных по программе лояльности",
}

PRODUCTS = (
    Product("A", "Горшок", 700, 300, category="Горшки"),
    Product("B", "Бант", 350, 50, category="Декор"),
    # Нулевая себестоимость: доходность показывается как «Нет себестоимости».
    Product("C", "Подарочный набор", 0, 0, category="Декор"),
    # Только хранение без продаж: доходность показывается как «Нет продаж».
    Product("D", "Подставка", 150, 50),
    # Во второй неделе только возврат: продажи меньше нуля.
    Product("E", "Кашпо", 250, 50, category="Горшки"),
)
NM_ID = {"A": 1001, "B": 1002, "C": 1003, "D": 1004, "E": 1005}
NAMES = {product.article: product.name for product in PRODUCTS}
PLANNED_PRICES = {"A": 1_700.0, "B": 1_300.0}


def _row(
    document: str,
    reason: str,
    article: str,
    day: str,
    *,
    quantity: float = 0,
    country: str = "Россия",
    srid: str = "",
    **amounts: float,
) -> list[object]:
    values: list[object] = [0] * len(HEADERS)
    values[COLUMN["Предмет"]] = "Товар" if article else ""
    values[COLUMN["Код номенклатуры"]] = NM_ID.get(article, "")
    values[COLUMN["Артикул поставщика"]] = article
    values[COLUMN["Название"]] = NAMES.get(article, "")
    values[COLUMN["Тип документа"]] = document
    values[COLUMN["Обоснование для оплаты"]] = reason
    values[COLUMN["Дата заказа покупателем"]] = day
    values[COLUMN["Дата продажи"]] = day
    values[COLUMN["Кол-во"]] = quantity
    values[COLUMN["Виды логистики, штрафов и корректировок ВВ"]] = ""
    values[COLUMN["Страна"]] = country
    values[COLUMN["Srid"]] = srid
    values[COLUMN["Коэффициент логистики"]] = 1
    for field, amount in amounts.items():
        values[COLUMN[FIELD_COLUMNS[field]]] = amount
    return values


def _sale(article: str, day: str, quantity: float, *, country: str = "Россия", **amounts: float):
    return _row("Продажа", "Продажа", article, day, quantity=quantity, country=country,
                srid=f"s-{article}-{day}", **amounts)


def _return(article: str, day: str, quantity: float, **amounts: float):
    return _row("Возврат", "Возврат", article, day, quantity=quantity,
                srid=f"r-{article}-{day}", **amounts)


def _service(reason: str, article: str, day: str, *, country: str = "Россия", **amounts: float):
    return _row("", reason, article, day, country=country, **amounts)


# Неделя 1: основной отчёт (MAIN), отчёт по выкупам и уведомление (BUYOUT).
WEEK1_MAIN = (
    _sale("A", "08.09.2026", 2, retail=3_000, realized=2_400, acquiring=48, commission=300,
          commission_vat=60, payout=1_920, loyalty_fee=30, loyalty_points=20),
    _service("Логистика", "A", "08.09.2026", logistics=150),
    _sale("B", "09.09.2026", 1, retail=1_200, realized=1_000, acquiring=20, commission=125,
          commission_vat=25, payout=800),
    _service("Штраф", "B", "09.09.2026", penalty=100),
    _sale("C", "10.09.2026", 1, retail=700, realized=600, acquiring=12, commission=75,
          commission_vat=15, payout=480),
    _service("Хранение", "D", "11.09.2026", storage=45),
    _sale("E", "11.09.2026", 1, retail=900, realized=700, acquiring=14, commission=87.5,
          commission_vat=17.5, payout=560),
    _service("Хранение", "", "12.09.2026", storage=300),
    _service("Удержания", "", "12.09.2026", deductions=250),
    _service("Операции на приемке", "", "13.09.2026", acceptance=60),
)
WEEK1_BUYOUT = (
    _sale("A", "10.09.2026", 1, country="Казахстан", retail=1_500, realized=1_200, acquiring=24,
          commission=150, commission_vat=30, payout=960),
    _service("Логистика", "A", "10.09.2026", country="Казахстан", logistics=60),
)
# Неделя 2: только основной отчёт; возврат E и компенсация без артикула.
WEEK2_MAIN = (
    _sale("A", "15.09.2026", 1, retail=1_600, realized=1_300, acquiring=26, commission=162.5,
          commission_vat=32.5, payout=1_040),
    _service("Логистика", "A", "15.09.2026", logistics=80),
    _return("E", "16.09.2026", 1, retail=900, realized=700, acquiring=14, commission=87.5,
            commission_vat=17.5, payout=560),
    _sale("B", "17.09.2026", 2, retail=2_400, realized=2_000, acquiring=40, commission=250,
          commission_vat=50, payout=1_600, loyalty_compensation=50),
    _service("Удержания", "", "18.09.2026", deductions=199),
    _service("Компенсация ущерба", "", "19.09.2026", payout=150),
)

TOTAL_KEYS = (
    "units", "revenue", "realized_revenue", "seller_payout", "financial_result",
    "cost_sold", "tax", "net_profit", "unallocated",
)
PRODUCT_KEYS = (
    "units", "main_units", "buyout_units", "main_revenue", "buyout_revenue", "revenue",
    "taxable_income", "seller_payout", "financial_result", "cost_sold", "tax",
    "net_profit", "wb_commission", "logistics_cost", "loyalty_cost",
)
PRODUCT_RATIO_KEYS = ("profitability", "commission_share", "logistics_share", "points_share", "net_margin")
SCENARIO_KEYS = (
    "current_price", "planned_price", "price_change", "profitability", "planned_revenue",
    "commission_rate", "planned_commission", "planned_points", "taxable_base", "tax",
    "profit", "net_profit_per_unit", "net_profit_total",
)
RATIO_NAMES = set(PRODUCT_RATIO_KEYS) | {
    "tax_rate", "product_profitability", "activity_profitability", "price_change",
    "commission_rate", "commission_share", "logistics_share", "points_share",
}


def _write_weekly(path: Path, rows) -> None:
    workbook = Workbook()
    sheet = workbook.active
    sheet.append(HEADERS)
    for row in rows:
        sheet.append(list(row))
    workbook.save(path)
    workbook.close()


def _write_notice(path: Path, number: str, notice_date: str, rows) -> None:
    workbook = Workbook()
    sheet = workbook.active
    sheet["A3"] = f"УВЕДОМЛЕНИЕ О ВЫКУПЕ №{number} от {notice_date}"
    for _ in range(6):
        sheet.append([])
    sheet.append(["№ п/п", "Артикул", "Наименование", "Количество", "Сумма выкупа, руб., (вкл. НДС)"])
    total_quantity = total_amount = 0.0
    for index, (article, quantity, amount) in enumerate(rows, start=1):
        sheet.append([index, article, NAMES[article], quantity, amount])
        total_quantity += quantity
        total_amount += amount
    sheet.append(["Итого:", None, None, total_quantity, total_amount])
    workbook.save(path)
    workbook.close()


def write_synthetic_reports(root: Path) -> list[Path]:
    """Two synthetic WB weeks; the first one has MAIN and BUYOUT channels."""
    root.mkdir(parents=True, exist_ok=True)
    week1 = root / "Еженедельный детализированный отчет №101.xlsx"
    buyout = root / "Еженедельный детализированный отчет №102.xlsx"
    notice = root / "Уведомление о выкупе №102 от 2026-09-10.xlsx"
    week2 = root / "Еженедельный детализированный отчет №103.xlsx"
    _write_weekly(week1, WEEK1_MAIN)
    _write_weekly(buyout, WEEK1_BUYOUT)
    # Договорная цена выкупа = перечисление по продаже − логистика выкупа.
    _write_notice(notice, "102", "2026-09-10", [("A", 1, 900)])
    _write_weekly(week2, WEEK2_MAIN)
    return [week1, buyout, notice, week2]


def import_synthetic_reports(root: Path) -> AppService:
    service = AppService(root / "data")
    for product in PRODUCTS:
        service.db.save_product(product, source="Эталон")
    service.db.set_setting("tax_rate", TAX_RATE)
    batch = service.prepare_import(write_synthetic_reports(root / "reports"))
    service.complete_import_batch(batch.sessions)
    return service


def _number(value) -> float | None:
    return None if value is None else float(value)


def _calculation_metrics(calculation: RunCalculation) -> dict[str, object]:
    totals = calculation.totals()
    rate = calculation.tax_rate
    cost_sold = totals["cost_sold"]
    product_net = sum(item.net_profit(rate) for item in calculation.products)
    products = {}
    for result in calculation.products:
        products[result.article] = {
            "units": result.units,
            "main_units": result.main_units_total,
            "buyout_units": result.buyout_units_total,
            "main_revenue": result.main_revenue_total,
            "buyout_revenue": result.buyout_revenue_total,
            "revenue": result.revenue_including_points,
            "taxable_income": result.taxable_income,
            "seller_payout": result.seller_payout,
            "financial_result": result.financial_result,
            "cost_sold": result.cost_sold,
            "tax": result.tax(rate),
            "net_profit": result.net_profit(rate),
            "wb_commission": result.wb_commission,
            "logistics_cost": result.logistics_cost,
            "loyalty_cost": result.loyalty_cost,
            "profitability": result.profitability(rate),
            "commission_share": result.commission_share(),
            "logistics_share": result.logistics_share(),
            "points_share": result.points_share(),
            "net_margin": result.net_margin(rate),
        }
    return {
        "tax_rate": rate,
        "totals": {key: totals[key] for key in TOTAL_KEYS},
        "product_net_profit": product_net,
        "activity_net_profit": totals["net_profit"],
        "product_profitability": product_net / cost_sold if cost_sold else 0.0,
        "activity_profitability": totals["net_profit"] / cost_sold if cost_sold else 0.0,
        "revenue_shares": calculation.revenue_shares(),
        "revenue_amounts": calculation.revenue_amounts(),
        "products": products,
    }


def _scenario_metrics(calculation: RunCalculation, prices: dict[str, float]) -> dict[str, object]:
    rows = {}
    for result in calculation.products:
        row = calculate_scenario(result, calculation.tax_rate, prices.get(result.article))
        rows[result.article] = {key: _number(getattr(row, key)) for key in SCENARIO_KEYS}
    return rows


def _export_cells(calculation: RunCalculation, destination: Path, prices=None, control=None) -> dict[str, object]:
    export_calculation(calculation, destination, prices, source_control_total=control)
    workbook = load_workbook(destination)
    try:
        cells: dict[str, object] = {}
        for sheet in workbook.worksheets:
            for row in sheet.iter_rows():
                for cell in row:
                    if cell.value is not None:
                        cells[f"{sheet.title}!{cell.coordinate}"] = cell.value
        return cells
    finally:
        workbook.close()


def build_snapshot(root: Path) -> dict[str, object]:
    """Build every etalon value from synthetic reports through the real pipeline."""
    with patch.dict(os.environ, {"WB_APP_DATA": str(root / "app")}):
        service = import_synthetic_reports(root)
        runs = service.db.list_runs()
        calculations = {
            f"{run.period_start}..{run.period_end}": service.db.load_calculation(run.id)
            for run in runs
        }
        first_run = runs[0]
        for article, price in PLANNED_PRICES.items():
            service.db.save_planned_price(first_run.id, article, price)
        overview = aggregate_calculations(list(calculations.values()))
        snapshot: dict[str, object] = {
            "reports": {key: _calculation_metrics(value) for key, value in calculations.items()},
            "overview": _calculation_metrics(overview),
            "scenario": _scenario_metrics(
                service.db.load_calculation(first_run.id),
                service.db.planned_prices(first_run.id),
            ),
            "export": {},
        }
        exports_dir = root / "exports"
        for run in runs:
            key = f"{run.period_start}..{run.period_end}"
            control = sum(float(row["total_amount"]) for row in service.db.list_source_files(run.id))
            snapshot["export"][key] = _export_cells(
                calculations[key],
                exports_dir / f"{key}.xlsx",
                service.db.planned_prices(run.id),
                control,
            )
        snapshot["export"]["overview"] = _export_cells(overview, exports_dir / "overview.xlsx")
        return snapshot


class CalculationEtalonTests(unittest.TestCase):
    """До и после интерфейсных изменений все показатели совпадают до копейки."""

    @classmethod
    def setUpClass(cls) -> None:
        cls.expected = json.loads(FIXTURE.read_text(encoding="utf-8"))
        directory = tempfile.TemporaryDirectory()
        cls.addClassCleanup(directory.cleanup)
        cls.actual = build_snapshot(Path(directory.name))

    def assert_values_match(self, expected, actual, path: str) -> None:
        if isinstance(expected, dict):
            self.assertIsInstance(actual, dict, path)
            self.assertEqual(sorted(actual), sorted(expected), path)
            for key in expected:
                self.assert_values_match(expected[key], actual[key], f"{path}.{key}")
            return
        if expected is None or isinstance(expected, str):
            self.assertEqual(actual, expected, path)
            return
        self.assertIsNotNone(actual, path)
        leaf = path.rsplit(".", 1)[-1]
        ratio = leaf in RATIO_NAMES or leaf.endswith("_share") or leaf == "net_margin"
        tolerance = RATIO_TOLERANCE if ratio else MONEY_TOLERANCE
        self.assertLessEqual(abs(float(actual) - float(expected)), tolerance, path)

    def test_fixture_covers_every_channel_and_edge_case(self) -> None:
        week1 = self.expected["reports"]["2026-09-08..2026-09-13"]["products"]
        week2 = self.expected["reports"]["2026-09-15..2026-09-19"]["products"]
        self.assertGreater(week1["A"]["buyout_revenue"], 0)
        self.assertGreater(week1["A"]["main_revenue"], 0)
        self.assertEqual(week1["D"]["units"], 0)
        self.assertEqual(week1["C"]["cost_sold"], 0)
        self.assertLess(week2["E"]["units"], 0)
        self.assertNotEqual(self.expected["overview"]["totals"]["unallocated"], 0)

    def test_report_totals_taxes_profit_and_shares_match_etalon(self) -> None:
        for key in ("reports", "overview"):
            with self.subTest(section=key):
                self.assert_values_match(self.expected[key], self.actual[key], key)

    def test_tax_base_is_main_plus_buyout_revenue_without_unallocated_tax(self) -> None:
        for key, metrics in (*self.actual["reports"].items(), ("overview", self.actual["overview"])):
            with self.subTest(report=key):
                rate = metrics["tax_rate"]
                products = metrics["products"].values()
                base = sum(item["main_revenue"] + item["buyout_revenue"] for item in products)
                self.assertLessEqual(abs(metrics["totals"]["tax"] - base * rate), MONEY_TOLERANCE)
                # No tax is charged on unallocated operations.
                self.assertLessEqual(
                    abs(
                        metrics["activity_net_profit"]
                        - metrics["product_net_profit"]
                        - metrics["totals"]["unallocated"]
                    ),
                    MONEY_TOLERANCE,
                )

    def test_price_scenario_matches_etalon(self) -> None:
        self.assert_values_match(self.expected["scenario"], self.actual["scenario"], "scenario")

    def test_export_cells_match_etalon(self) -> None:
        for workbook_key, expected_cells in self.expected["export"].items():
            actual_cells = self.actual["export"][workbook_key]
            with self.subTest(workbook=workbook_key):
                self.assertEqual(sorted(actual_cells), sorted(expected_cells))
                for coordinate, expected in expected_cells.items():
                    actual = actual_cells[coordinate]
                    if isinstance(expected, str):
                        self.assertEqual(actual, expected, coordinate)
                    else:
                        self.assertLessEqual(
                            abs(float(actual) - float(expected)), MONEY_TOLERANCE, coordinate
                        )


def _write_fixture() -> None:
    with tempfile.TemporaryDirectory() as directory:
        snapshot = build_snapshot(Path(directory))
    FIXTURE.parent.mkdir(parents=True, exist_ok=True)
    FIXTURE.write_text(
        json.dumps(snapshot, ensure_ascii=False, indent=1, sort_keys=True) + "\n",
        encoding="utf-8",
    )
    print(f"Эталон записан: {FIXTURE}")


if __name__ == "__main__":
    if "--write" in sys.argv:
        _write_fixture()
    else:
        unittest.main()
